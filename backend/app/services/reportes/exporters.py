from io import BytesIO
from decimal import Decimal
from datetime import date, datetime, timezone
from html import escape
import re
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle


def safe_cell(value):
    if isinstance(value,str):
        value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',value)
        if len(value)>32767:
            raise ValueError('Una celda supera el límite de texto de Excel. Acota el reporte o revisa ese registro.')
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def typed(value, type_name):
    if value is None:
        return None
    if type_name == 'number':
        return Decimal(str(value))
    if type_name == 'date':
        return date.fromisoformat(str(value))
    if type_name == 'datetime':
        stamp = datetime.fromisoformat(str(value))
        return stamp.astimezone(timezone.utc).replace(tzinfo=None) if stamp.tzinfo else stamp
    return safe_cell(str(value))


def export(result, format):
    out = BytesIO()
    columns = result['columnas']
    presentation = result['solicitud'].get('presentacion', {})
    if format == 'xlsx':
        # Validar antes de abrir hojas temporales: un rechazo no deja generators abiertos.
        for row in result['filas']:
            for column in columns:
                typed(row.get(column['campo']),column['tipo'])
        for value in list(result['resumen'].values())+result['advertencias']+list(result['solicitud']['filtros'].values())+result.get('recomendaciones',[]):
            safe_cell(str(value))
        book = Workbook(write_only=True)
        summary = book.create_sheet('Resumen')
        summary.append(['Reporte', safe_cell(result['titulo'])])
        summary.append(['Fecha de corte', result['corte']])
        summary.append(['Empresa', result['id_empresa']])
        summary.append(['Fechas y horas', 'Las celdas de fecha/hora se expresan en UTC.'])
        for k, v in result['resumen'].items():
            summary.append([safe_cell(k), safe_cell(str(v))])
        for note in result['advertencias']:
            summary.append(['Nota', safe_cell(note)])
        for recommendation in result.get('recomendaciones',[]):
            summary.append(['Recomendación',safe_cell(str(recommendation))])
        filters = book.create_sheet('Filtros')
        for k, v in result['solicitud']['filtros'].items():
            filters.append([safe_cell(k), safe_cell(str(v))])
        for k, v in presentation.items():
            filters.append([safe_cell('presentacion.'+k),safe_cell(str(v))])
        detail = book.create_sheet('Detalle')
        detail.freeze_panes = 'A2'
        detail.append([c['titulo'] for c in columns])
        for row in result['filas']:
            detail.append([typed(row.get(c['campo']), c['tipo']) for c in columns])
        book.save(out)
    elif format == 'pdf':
        if any(len(str(row.get(c['campo'])))>2000 for row in result['filas'] for c in columns):
            raise ValueError('Hay textos demasiado extensos para PDF. Exporta el detalle a Excel o acota el reporte.')
        styles = getSampleStyleSheet()
        small = styles['BodyText']
        small.fontSize = 6
        small.leading = 8
        para = lambda value: Paragraph(escape(str(value if value is not None else '—')), small)
        paper = A4 if presentation.get('orientacion')=='vertical' else landscape(A4)
        doc = SimpleDocTemplate(out, pagesize=paper, rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24)
        story = [Paragraph(escape(result['titulo']), styles['Title']),
                 para(f"Empresa {result['id_empresa']} · corte {result['corte']} · {len(result['filas'])} registros"), Spacer(1,10)]
        for k, v in result['resumen'].items():
            story.append(para(f'{k}: {v}'))
        for note in result['advertencias']:
            story.append(para(note))
        for recommendation in result.get('recomendaciones',[]):
            story.append(para('Recomendación: '+str(recommendation)))
        story.extend([para('Filtros: ' + str(result['solicitud']['filtros'])), Spacer(1,10)])
        # Dividir columnas en bandas preserva todos los campos sin texto ilegible.
        band_size = 5 if presentation.get('orientacion')=='vertical' else 8
        for start in range(0, len(columns), band_size):
            band = columns[start:start+band_size]
            data = [[para(c['titulo']) for c in band]] + [[para(row.get(c['campo'])) for c in band] for row in result['filas']]
            table = LongTable(data,repeatRows=1,colWidths=[(paper[0]-48)/max(len(band),1)]*len(band))
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e2e8f0')),('GRID',(0,0),(-1,-1),0.25,colors.lightgrey),('VALIGN',(0,0),(-1,-1),'TOP')]))
            story.extend([table, Spacer(1,10)])
        doc.build(story)
    else:
        raise ValueError('Formato no soportado.')
    return out.getvalue()
