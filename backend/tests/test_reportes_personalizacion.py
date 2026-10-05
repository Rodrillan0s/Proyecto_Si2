"""Regresiones de contratos, fechas, análisis financiero y asistente único, sin red."""
import unittest
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch, MagicMock
from fastapi import HTTPException
from openpyxl import load_workbook
from app.services import reportes_services as reports
from app.services.ai import assistant_service as assistant
from app.services.reportes.catalog import REGISTRY
from app.services.reportes.contracts import ReportRequest, Interpretation
from app.services.reportes.authorization import allowed
from app.services.reportes.exporters import export
from app.services.reportes.scheduling import period_request
from app.services.ai.providers import ProviderError
from test_reportes import actor, fake_transaction


class CustomReportsTests(unittest.TestCase):
    def test_scheduled_period_preserves_presentation(self):
        configuration = {'solicitud': {'reporte':'incidencias','filtros':{},
            'presentacion':{'titulo':'Control semanal','columnas':['titulo'],'orientacion':'vertical'}},
            'zona':'America/La_Paz','periodo':'mes_anterior'}
        result = period_request(configuration,datetime(2026,10,4,tzinfo=timezone.utc))
        self.assertEqual(result['presentacion'],configuration['solicitud']['presentacion'])
        self.assertEqual(result['filtros']['hasta'],'2026-09-30')
    def build(self, request, data):
        with patch.object(reports,'authorize',return_value=(REGISTRY[request.reporte],[9])), \
             patch.object(reports.repo,'datasets',return_value=data), \
             patch.object(reports.repo,'cutoff',return_value={'corte':datetime(2026,10,4,tzinfo=timezone.utc)}):
            return reports.build(None,actor(),request)[0]

    def test_selected_columns_sort_null_last_and_export_match(self):
        request = ReportRequest(reporte='incidencias', presentacion={
            'titulo':'=SUM(1,2)', 'columnas':['titulo','id_incidencia'],
            'ordenar_por':'id_incidencia','orden':'desc','orientacion':'vertical'})
        data = [dict(id_incidencia=i,titulo=str(i),prioridad='MEDIA') for i in (None,1,10)]
        result = self.build(request,data)
        self.assertEqual([r['id_incidencia'] for r in result['filas']],[10,1,None])
        self.assertEqual([c['campo'] for c in result['columnas']],['titulo','id_incidencia'])
        book = load_workbook(BytesIO(export(result,'xlsx')))
        self.assertEqual(book['Detalle'].max_column,2)
        self.assertEqual(book['Detalle']['B2'].value,10)
        self.assertEqual(book['Resumen']['B1'].value,"'=SUM(1,2)")
        self.assertEqual(book['Detalle'].freeze_panes,'A2')
        self.assertTrue(export(result,'pdf').startswith(b'%PDF'))
        # Campos no visibles permanecen para intersectar filas en envíos.
        self.assertIn('prioridad',result['filas'][0])

    def test_invalid_and_duplicate_columns_are_rejected_even_empty(self):
        for columns in [['password'],['titulo','titulo']]:
            with self.assertRaises(HTTPException) as error:
                self.build(ReportRequest(reporte='incidencias',presentacion={'columnas':columns}),[])
            self.assertEqual(error.exception.status_code,422)
        empty = self.build(ReportRequest(reporte='incidencias',presentacion={'columnas':['titulo']}),[])
        self.assertEqual(empty['columnas'][0]['campo'],'titulo')

    def test_period_totals_are_decimal_and_keep_currency_groups(self):
        result = self.build(ReportRequest(reporte='costos_periodo'),[
            {'id_obra':9,'obra':'A','moneda':'BOB','monto':Decimal('.10')},
            {'id_obra':9,'obra':'A','moneda':'BOB','monto':Decimal('.20')},
            {'id_obra':10,'obra':'B','moneda':'USD','monto':Decimal('5')}])
        self.assertEqual([r['monto_registrado'] for r in result['resumen']['costos_por_obra']],[Decimal('.30'),Decimal('5')])

    def test_date_parameters_are_inclusive_and_bound(self):
        for kind in ('presupuestario','asignacion_personal','utilizacion_personal','avance_ejecutivo','costos_periodo'):
            with self.subTest(kind=kind), patch.object(reports.repo,'rows',return_value=[]) as query:
                reports.repo.datasets(None,actor(),ReportRequest(reporte=kind,filtros={'desde':'2026-01-01','hasta':'2026-10-04'}),[9])
                sql,params = query.call_args.args[1:]
                self.assertIn('>= %s',sql)
                self.assertIn('<= %s',sql)
                self.assertIn(date(2026,1,1),params)
                if kind=='avance_ejecutivo':
                    self.assertLess(sql.index('av.fecha_registro >='),sql.index('ORDER BY av.fecha_registro'))

    def test_finance_requires_both_permissions_and_excludes_clients(self):
        definition = REGISTRY['saldo_comercial']
        for role,perms in [('ADMINISTRADOR_EMPRESA',['Visualizar_control_costos']),
                           ('ADMINISTRADOR_EMPRESA',['Visualizar_clientes']),
                           ('CLIENTE',['Visualizar_control_costos','Visualizar_clientes'])]:
            self.assertFalse(allowed(actor(role,perms),definition))
        self.assertTrue(allowed(actor(permissions=['Visualizar_clientes','Visualizar_control_costos']),definition))

    def test_financial_rows_require_sales_amount_and_costs(self):
        rows = [dict(id_obra=9,obra='A',moneda='BOB',unidades_vendidas=1,ventas_sin_monto=0,monto_pactado=Decimal('100.05'),costos_registrados=Decimal('20.03')),
                dict(id_obra=10,obra='B',moneda='USD',unidades_vendidas=1,ventas_sin_monto=1,monto_pactado=None,costos_registrados=Decimal('20')),
                dict(id_obra=11,obra='C',moneda='BOB',unidades_vendidas=0,ventas_sin_monto=0,monto_pactado=None,costos_registrados=None)]
        with patch.object(reports.repo,'rows',return_value=rows) as query:
            data = reports.repo.datasets(None,actor(),ReportRequest(reporte='saldo_comercial'),[9,10,11])
        self.assertEqual(data[0]['saldo_estimado'],Decimal('80.02'))
        self.assertIsNone(data[1]['saldo_estimado'])
        self.assertIsNone(data[2]['saldo_estimado'])
        sql = query.call_args.args[1]
        self.assertIn('DISTINCT ON (cu.id_unidad)',sql)
        self.assertIn('cli.id_empresa=o.id_empresa',sql)
        self.assertIn("ce.estado='REGISTRADO'",sql)


class UnifiedAssistantTests(unittest.TestCase):
    def test_profit_answer_uses_backend_evidence_without_calling_llm(self):
        result = {'reporte':'saldo_comercial','solicitud':{'reporte':'saldo_comercial','filtros':{}},
            'titulo':'Saldo comercial', 'resumen':{'obras_sin_datos_suficientes':1},
            'filas':[{'obra':'Alameda','saldo_estimado':'80.02','moneda':'BOB','estado':'SALDO_POSITIVO'}]}
        finance_actor=actor(permissions=['Visualizar_reportes','Visualizar_clientes','Visualizar_control_costos'])
        with patch.object(reports.repo,'transaction',fake_transaction),patch.object(assistant,'context',return_value=finance_actor), \
             patch.object(reports.repo,'works',return_value=[]),patch.object(reports,'create',return_value={'id':'e','resultado':result}), \
             patch.object(reports,'owned'),patch.object(assistant,'get_provider') as provider:
            response=assistant.consult({'nro_usuario':1},Interpretation(texto='¿Qué obras generaron ganancias?'),7)
        self.assertIn('Alameda: 80.02 BOB',response['respuesta'])
        self.assertIn('no cobros',response['respuesta'])
        self.assertEqual(response['tema'],'finanzas')
        provider.assert_not_called()

    def test_natural_cost_period_uses_existing_dated_query(self):
        from app.services.reportes.interpreter import interpret
        result=interpret('costos este mes',list(REGISTRY),[])
        self.assertEqual(result['estado'],'ready')
        self.assertEqual(result['solicitud']['reporte'],'costos_periodo')
    def test_profit_question_and_typo_use_financial_capability(self):
        for question in ('¿Cuáles son las obras que han generado ganancias? O aún no hay ni una?', '¿Hay ganacias?', '¿Qué obra ha generado ganancias?', '¿Qué obra tiene pérdidas?'):
            result = assistant._plan(question,list(REGISTRY),[],None,None,True)
            self.assertEqual(result['solicitud']['reporte'],'saldo_comercial')

    def test_units_are_not_confused_with_utilities(self):
        result = assistant._plan('estado de unidades',list(REGISTRY),[],None,None,True)
        self.assertEqual(result['solicitud']['reporte'],'estado_unidades')

    def test_commercial_and_followup_share_same_planner(self):
        self.assertEqual(assistant._plan('prospectos',[],[],None,None,True).accion,'crm')
        self.assertEqual(assistant._plan('¿Y cuáles están interesados?',[],[],None,'crm',True).accion,'crm')

    def test_untrusted_plan_cannot_inject_sql_or_unauthorized_report(self):
        provider = MagicMock()
        for raw in ('{"accion":"crm"}', '{"accion":"reportes","solicitudes":[{"reporte":"saldo_comercial"}]}',
                    '{"accion":"reportes","sql":"SELECT *","solicitudes":[]}'):
            provider.complete.return_value=raw
            with patch.object(assistant,'get_provider',return_value=provider):
                result = assistant._plan('informe especial',['incidencias'],[],None,None,False)
                self.assertEqual(result['estado'],'unsupported')

    def test_foreign_conversation_rejected_before_provider(self):
        with patch.object(reports.repo,'transaction',fake_transaction),patch.object(assistant,'context',return_value=actor()), \
             patch.object(reports.repo,'works',return_value=[]),patch.object(reports.repo,'get_conversation',return_value=None), \
             patch.object(assistant,'get_provider') as provider:
            with self.assertRaises(HTTPException) as error:
                assistant.consult({'nro_usuario':1},Interpretation(texto='ventas',conversacion='00000000-0000-4000-8000-000000000001'),7)
            self.assertEqual(error.exception.status_code,404)
            provider.assert_not_called()


if __name__=='__main__':
    unittest.main()
