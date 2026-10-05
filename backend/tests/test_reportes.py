"""Pruebas aisladas: no conectan DB, no invocan IA ni envían correos reales."""
import unittest
from contextlib import contextmanager
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch
import requests
from fastapi import HTTPException
from openpyxl import load_workbook
from app.services.reportes.authorization import Context, context, authorize, allowed
from app.services.reportes.catalog import REGISTRY
from app.services.reportes.contracts import ReportRequest, Filters, Interpretation
from app.services.reportes.interpreter import interpret
from app.services.reportes.exporters import export
from app.services.reportes.scheduling import next_run, period_request
from app.services.reportes import worker
from app.services import reportes_services as services, email_service
from app.services.ai.providers import DeepSeekProvider, ProviderError, get_provider
from app.config import Config


def actor(role='ADMINISTRADOR_EMPRESA', permissions=None, uid=1, company=7):
    return Context(uid,company,role,10,frozenset(permissions or ['Visualizar_reportes','Exportar_reportes','Enviar_reportes','Visualizar_incidencias','Visualizar_inventario','Visualizar_avances']), 'prueba@example.test')


@contextmanager
def fake_transaction(*args,**kwargs):
    yield MagicMock()


class AuthorizationTests(unittest.TestCase):
    def test_tenant_is_reloaded_not_taken_from_jwt(self):
        row = dict(id_usuario=1,id_empresa=7,id_persona=10,estado='ACTIVO',correo='a@example.test',nombre_rol='CLIENTE',permisos=['Visualizar_avances'])
        with patch.object(services.repo,'actor',return_value=row), patch.object(services.repo,'one',return_value={'id_empresa':7}):
            with self.assertRaises(HTTPException) as err: context(None,1,8)
            self.assertEqual(err.exception.status_code,403)
            self.assertEqual(context(None,1).empresa,7)
            row['estado']='INACTIVO'
            with self.assertRaises(HTTPException): context(None,1,7)

    def test_global_requires_explicit_tenant(self):
        row = dict(id_usuario=1,id_empresa=1,id_persona=10,estado='ACTIVO',correo='a@example.test',nombre_rol='ADMINISTRADOR',permisos=[])
        with patch.object(services.repo,'actor',return_value=row):
            with self.assertRaises(HTTPException) as err: context(None,1)
            self.assertEqual(err.exception.status_code,422)

    def test_client_cannot_query_costs_or_inventory_even_with_domain_grant(self):
        ctx = actor('CLIENTE', ['Visualizar_reportes','Visualizar_control_costos','Visualizar_inventario','Visualizar_avances'])
        self.assertFalse(allowed(ctx,REGISTRY['comparativo_costos']))
        self.assertFalse(allowed(ctx,REGISTRY['stock']))
        self.assertTrue(allowed(ctx,REGISTRY['avance_ejecutivo']))

    def test_reject_cross_work_and_unsupported_filters(self):
        with patch.object(services.repo,'works',return_value=[{'id_obra':9}]):
            with self.assertRaises(HTTPException): authorize(None,actor(),ReportRequest(reporte='incidencias',filtros=Filters(id_obra=99)))
            with self.assertRaises(HTTPException): authorize(None,actor(),ReportRequest(reporte='stock',filtros=Filters(desde='2026-01-01')))

    def test_download_reloads_role_and_resource_scope(self):
        execution = dict(id_usuario=1,id_empresa=7,solicitud={'reporte':'incidencias','filtros':{}},politica=actor().policy(),alcance=[9])
        with patch.object(services.repo,'one',return_value=execution), patch.object(services,'context',return_value=actor()), patch.object(services,'authorize',return_value=(REGISTRY['incidencias'],[])):
            with self.assertRaises(HTTPException): services.owned(None,'x',{'nro_usuario':1})

    def test_download_other_actor_is_forbidden(self):
        with patch.object(services.repo,'one',return_value={'id_usuario':2,'id_empresa':7}), patch.object(services,'context',return_value=actor()):
            with self.assertRaises(HTTPException): services.owned(None,'x',{'nro_usuario':1})

    def test_cliente_scope_uses_persona_fk_never_email(self):
        with patch.object(services.repo,'rows',return_value=[]) as query:
            services.repo.works(None,actor('CLIENTE'))
            sql,params = query.call_args.args[1:]
            self.assertIn('d.id_cliente=%s',sql)
            self.assertEqual(params,[7,10])
            self.assertNotIn('correo',sql)


class LanguageTests(unittest.TestCase):
    def setUp(self):
        self.catalog = list(REGISTRY)
        self.works = [{'id_obra':9,'nombre':'Los Pinos','codigo':'P-9'},{'id_obra':10,'nombre':'Los Pinos Norte','codigo':'P-10'}]
        self.now = datetime(2026,10,4,tzinfo=timezone.utc)

    def test_typos_accents_and_critical_specificity(self):
        for text in ['reporte de stok','muéstrame el inventaro','existencias']:
            self.assertEqual(interpret(text,self.catalog,self.works)['solicitud']['reporte'],'stock')
        response = interpret('incidencias críticas del mes pasado',self.catalog,self.works,now=self.now)
        self.assertEqual(response['solicitud'],{'reporte':'incidencias_criticas','filtros':{'desde':'2026-09-01','hasta':'2026-09-30'}})

    def test_ambiguity_multiple_reports_unknown_work_and_continuation(self):
        self.assertEqual(interpret('stock o costos',self.catalog,self.works)['estado'],'needs_clarification')
        self.assertEqual(interpret('incidencias proyecto desconocido',self.catalog,self.works)['estado'],'needs_clarification')
        self.assertEqual(interpret('incidencias Los Pinos Norte',self.catalog,self.works)['estado'],'needs_clarification')
        response = interpret('ahora en PDF',self.catalog,self.works,{'reporte':'incidencias','filtros':{'id_obra':9}})
        self.assertEqual(response['solicitud']['filtros']['id_obra'],9)
        self.assertEqual(response['formato'],'pdf')

    def test_near_project_and_missing_baseline_work(self):
        response = interpret('incidencias proyecto Los Pnios',self.catalog,[self.works[0]])
        self.assertEqual(response['solicitud']['filtros']['id_obra'],9)
        self.assertEqual(interpret('comparativo de costos',self.catalog,self.works)['estado'],'needs_clarification')

    def test_state_reports_do_not_silently_ignore_dates(self):
        self.assertEqual(interpret('stock este mes',self.catalog,[],now=self.now)['estado'],'needs_clarification')
        with self.assertRaises(ValueError): ReportRequest(reporte='incidencias',filtros={'desde':'2026-10-04','hasta':'2026-09-01'})
        with self.assertRaises(ValueError): ReportRequest(reporte='incidencias',sql='select * from users')

    def test_priority_stock_state_and_category_are_preserved(self):
        response = interpret('stock bajo categoria 3',self.catalog,[])
        self.assertEqual(response['solicitud']['filtros'],{'estado':'STOCK_BAJO','id_categoria':3})
        response = interpret('incidencias de prioridad alta resueltas',self.catalog,[])
        self.assertEqual(response['solicitud']['filtros'],{'prioridad':'ALTA','estado':'RESUELTA'})


class ExportTests(unittest.TestCase):
    def result(self):
        return dict(titulo='Costos <privados>',corte='2026-10-04T12:00:00Z',id_empresa=7,
            resumen={'total':'0.30'},advertencias=['<script> & nota'],solicitud={'filtros':{}},
            columnas=[dict(campo='concepto',titulo='Concepto',tipo='text'),dict(campo='monto',titulo='Monto',tipo='number')],
            filas=[{'concepto':'=HYPERLINK("https://example.test")','monto':'0.10'},{'concepto':'@SUM(1)','monto':'0.20'}])

    def test_xlsx_exact_rows_typed_numbers_formula_neutralized(self):
        book = load_workbook(BytesIO(export(self.result(),'xlsx')))
        sheet = book['Detalle']
        self.assertEqual(sheet.max_row,3)
        self.assertEqual(sheet['A2'].data_type,'s')
        self.assertTrue(sheet['A2'].value.startswith("'="))
        self.assertEqual(Decimal(str(sheet['B2'].value))+Decimal(str(sheet['B3'].value)),Decimal('.3'))
        self.assertEqual(sheet['B2'].data_type,'n')

    def test_pdf_markup_and_empty_result(self):
        result = self.result()
        self.assertTrue(export(result,'pdf').startswith(b'%PDF'))
        result['filas']=[];result['columnas']=[]
        self.assertTrue(export(result,'pdf').startswith(b'%PDF'))
        self.assertTrue(export(result,'xlsx').startswith(b'PK'))

    def test_comparison_decimal_totals_and_no_false_savings(self):
        rows = [dict(costo_presupuestado=Decimal('.10'),presupuesto_revisado=Decimal('.10'),costo_ejecutado=Decimal('.20'),variacion_revisada=Decimal('.10'),moneda='BOB')]*3
        with patch.object(services,'authorize',return_value=(REGISTRY['comparativo_costos'],[9])),patch.object(services.repo,'datasets',return_value=rows),patch.object(services.repo,'one',return_value={'corte':datetime.now(timezone.utc)}):
            result,_ = services.build(None,actor(),ReportRequest(reporte='comparativo_costos',filtros={'id_obra':9}))
            self.assertEqual(result['resumen']['costo_ejecutado'],Decimal('.60'))
            self.assertEqual(result['resumen']['variacion_revisada'],Decimal('.30'))
            self.assertTrue(result['recomendaciones'])

    def test_replenishment_subtracts_pending_and_requires_purchase_permission(self):
        stock = [dict(id_material=1,nombre_material='Cemento',stock_actual=Decimal('2'),stock_minimo=Decimal('10'))]
        ctx = actor(permissions=['Visualizar_reportes','Visualizar_inventario','Visualizar_ordenes_compra'])
        with patch.object(services,'authorize',return_value=(REGISTRY['stock'],[])),patch.object(services.repo,'datasets',return_value=stock),patch.object(services.repo,'cutoff',return_value={'corte':datetime.now(timezone.utc)}),patch.object(services.repo,'pending_purchases',return_value=[{'id_material':1,'pendiente':Decimal('3')}]) as pending:
            result,_ = services.build(None,ctx,ReportRequest(reporte='stock'))
            self.assertEqual(result['recomendaciones'][0]['reposicion_minima_sugerida'],'5')
            self.assertEqual(result['permisos_fuente'],['Visualizar_ordenes_compra'])
            pending.reset_mock()
            result,_ = services.build(None,ctx,ReportRequest(reporte='stock'),source_allowlist=[])
            self.assertNotIn('pendiente_recepcion',result['recomendaciones'][0])
            pending.assert_not_called()

    def test_source_permission_revocation_invalidates_saved_export(self):
        item = dict(id_usuario=1,id_empresa=7,solicitud={'reporte':'stock','filtros':{}},politica=actor().policy(),alcance=[],resultado={'permisos_fuente':['Visualizar_ordenes_compra']})
        with patch.object(services.repo,'get_execution',return_value=item),patch.object(services,'context',return_value=actor()),patch.object(services,'authorize',return_value=(REGISTRY['stock'],[])):
            with self.assertRaises(HTTPException): services.owned(None,'x',{'nro_usuario':1},'Exportar_reportes')

    def test_excel_dates_are_typed_and_text_limits_are_explicit(self):
        result = self.result()
        result['columnas']=[dict(campo='fecha',titulo='Fecha',tipo='datetime')]
        result['filas']=[{'fecha':'2026-10-04T08:00:00-04:00'}]
        book = load_workbook(BytesIO(export(result,'xlsx')))
        self.assertEqual(book['Detalle']['A2'].value,datetime(2026,10,4,12))
        result=self.result();result['filas'][0]['concepto']='x'*40000
        with self.assertRaises(ValueError): export(result,'xlsx')


class SchedulingTests(unittest.TestCase):
    def config(self,**kwargs):
        return dict(frecuencia='mensual',hora='08:00',zona='America/La_Paz',dia=31,periodo='mes_anterior',solicitud={'reporte':'incidencias','filtros':{}},**kwargs)

    def test_short_month_and_utc(self):
        due = next_run(self.config(),datetime(2026,2,1,tzinfo=timezone.utc))
        self.assertEqual(due,datetime(2026,2,28,12,tzinfo=timezone.utc))
        self.assertEqual(period_request(self.config(),due)['filtros'],{'desde':'2026-01-01','hasta':'2026-01-31'})

    def test_weekly_validation_and_unknown_zone(self):
        config = self.config(); config.update(frecuencia='semanal',dia=8)
        with self.assertRaises(HTTPException): next_run(config)
        config.update(dia=1,zona='bad/timezone')
        with self.assertRaises(HTTPException): next_run(config)

    def test_queue_claim_short_lock_and_no_reclaim_ambiguous_email(self):
        with patch.object(worker.repo,'transaction',fake_transaction),patch.object(worker.repo,'one',return_value=None) as one,patch.object(worker.repo,'rows',return_value=[]) as rows:
            worker.claim_delivery()
            self.assertIn('SKIP LOCKED',one.call_args.args[1])
            self.assertIn("estado='INCIERTO'",rows.call_args_list[0].args[1])
            worker.claim_execution()
            self.assertIn('SKIP LOCKED',one.call_args.args[1])
            self.assertIn("interval '20 minutes'",one.call_args.args[1])

    def test_deduplication_delivery_identity(self):
        with patch.object(services.repo,'one') as insert:
            services.enqueue_delivery(None,'execution',[1,1,2],['pdf'],'stable')
            self.assertEqual(insert.call_count,2)
            self.assertTrue(all('ON CONFLICT(identidad)' in c.args[1] for c in insert.call_args_list))

    def test_row_scope_intersection_prevents_worker_data_leak(self):
        sender = actor('ALBAÑIL',uid=1); recipient = actor(uid=2)
        item = dict(id_usuario=1,id_empresa=7,programacion=None,solicitud={'reporte':'incidencias','filtros':{}})
        with patch.object(worker,'context',return_value=sender),patch.object(services,'validate_recipient',return_value=recipient),patch.object(services,'build',side_effect=[({'filas':[{'id_incidencia':1}]},[9]),({'filas':[{'id_incidencia':1},{'id_incidencia':2}]},[9])]):
            with self.assertRaises(HTTPException): worker.recipient_result(None,item,{'id_usuario':2})

    def test_revoked_sender_permission_never_contacts_brevo(self):
        delivery = dict(id='job',ejecucion='exec',id_usuario=2,formatos=['pdf'],identidad='stable',intentos=1)
        execution = dict(id_usuario=1,id_empresa=7,programacion=None,solicitud={'reporte':'incidencias','filtros':{}})
        revoked = actor(permissions=['Visualizar_reportes','Visualizar_incidencias'])
        with patch.object(worker.repo,'transaction',fake_transaction),patch.object(worker.repo,'get_execution',return_value=execution),patch.object(worker,'context',return_value=revoked),patch.object(worker,'enviar_reporte') as email,patch.object(worker.repo,'complete_delivery') as record:
            worker.deliver(delivery)
            email.assert_not_called()
            self.assertEqual(record.call_args.args[1][0],'FALLIDO')

    def test_ambiguous_email_is_not_requeued(self):
        delivery = dict(id='job',ejecucion='exec',id_usuario=2,formatos=['pdf'],identidad='stable',intentos=1)
        execution = dict(id_usuario=1,id_empresa=7,programacion=None,solicitud={'reporte':'incidencias','filtros':{}})
        result = dict(reporte='incidencias',titulo='Incidencias',corte='2026-10-04',filas=[{'id_obra':9}],advertencias=[],permisos_fuente=[])
        with patch.object(worker.repo,'transaction',fake_transaction),patch.object(worker.repo,'get_execution',return_value=execution),patch.object(worker,'recipient_result',return_value=(result,actor(uid=2),actor())),patch.object(worker,'export',return_value=b'PDF'),patch.object(worker,'context',return_value=actor()),patch.object(services,'validate_recipient',return_value=actor(uid=2)),patch.object(worker,'authorize',return_value=(REGISTRY['incidencias'],[9])),patch.object(worker,'enviar_reporte',side_effect=email_service.ReportEmailError('Timeout',uncertain=True)),patch.object(worker.repo,'complete_delivery') as record:
            worker.deliver(delivery)
            self.assertEqual(record.call_args.args[1][0],'INCIERTO')


class TransportTests(unittest.TestCase):
    def test_assistant_reuses_report_and_survives_provider_outage(self):
        request = {'reporte':'stock','filtros':{}}
        from app.services.ai import assistant_service as assistant
        execution = {'id':'exec','resultado':{'reporte':'stock','solicitud':request,'titulo':'Stock','resumen':{'registros':3},'advertencias':['Existencias actuales.'],'recomendaciones':[], 'filas':[]}}
        with patch.object(services.repo,'transaction',fake_transaction),patch.object(assistant,'context',return_value=actor()),patch.object(services.repo,'works',return_value=[]),patch.object(services,'owned'),patch.object(services,'create',return_value=execution) as query,patch('app.services.ai.assistant_service.get_provider',side_effect=ProviderError()):
            result = services.assistant({'nro_usuario':1},Interpretation(texto='stock'),7)
            self.assertIn('3',result['respuesta'])
            self.assertEqual(result['ejecucion']['id'],'exec')
            self.assertEqual(query.call_args.args[1].reporte,'stock')

    def test_conversation_from_another_actor_or_company_is_rejected(self):
        with patch.object(services.repo,'transaction',fake_transaction),patch.object(services,'context',return_value=actor()),patch.object(services.repo,'works',return_value=[]),patch.object(services.repo,'get_conversation',return_value=None):
            with self.assertRaises(HTTPException) as error: services.interpretation({'nro_usuario':1},Interpretation(texto='ahora PDF',conversacion='00000000-0000-4000-8000-000000000001'),7)
            self.assertEqual(error.exception.status_code,404)

    def test_crm_contract_uses_injected_provider_and_fresh_authorization(self):
        from app.services.ai.ai_service import AIService
        crm_actor = actor(permissions=['Visualizar_clientes'])
        provider = MagicMock();provider.generar_respuesta.return_value='Hay 3 prospectos.'
        with patch.object(services.repo,'transaction',fake_transaction),patch('app.services.ai.ai_service.context',return_value=crm_actor),patch('app.services.ai.ai_service.ContextBuilder.construir_contexto_crm',return_value={'empresa_autorizada':7,'metricas_crm':{'prospectos':3}}),patch('app.services.ai.ai_service.crm_services._log'):
            response = AIService(provider).consultar_crm('Cuántos prospectos',{'nro_usuario':1},7)
            self.assertTrue(response['success'])
            self.assertEqual(response['respuesta'],'Hay 3 prospectos.')
            self.assertEqual(response['metricas'],{'prospectos':3})

    def test_server_speech_disabled_is_explicit(self):
        from app.services.reportes.speech import transcribe
        with patch.object(Config,'SPEECH_PROVIDER','disabled'):
            with self.assertRaises(HTTPException) as error: transcribe(b'audio')
            self.assertEqual(error.exception.status_code,503)

    def test_provider_maps_configuration_and_never_sql(self):
        response = MagicMock(status_code=200); response.json.return_value={'choices':[{'message':{'content':'{"reporte":"stock","filtros":{}}'}}]}
        with patch.object(Config,'DEEPSEEK_API_KEY','fake'),patch.object(Config,'DEEPSEEK_BASE_URL','https://api.example.test/v1'),patch.object(Config,'DEEPSEEK_MODEL','modelo-elegido'),patch('app.services.ai.providers.requests.post',return_value=response) as post:
            result = DeepSeekProvider().interpret_request('stock',[{'id':'stock'}],[])
            self.assertEqual(result['reporte'],'stock')
            self.assertEqual(post.call_args.kwargs['json']['model'],'modelo-elegido')
            self.assertEqual(post.call_args.args[0],'https://api.example.test/v1/chat/completions')
            self.assertNotIn('DB_PASSWORD',str(post.call_args))

    def test_provider_errors_do_not_expose_transport_body(self):
        response = MagicMock(status_code=401,text='secret detail')
        with patch.object(Config,'DEEPSEEK_API_KEY','fake'),patch('app.services.ai.providers.requests.post',return_value=response):
            with self.assertRaises(ProviderError) as exc: DeepSeekProvider().complete('system','prompt')
            self.assertNotIn('secret',str(exc.exception))

    def test_brevo_accepted_timeout_and_rate_limit(self):
        response = MagicMock(status_code=201);response.json.return_value={'messageId':'fake-id'}
        with patch.object(Config,'BREVO_API_KEY','fake'),patch.object(Config,'BREVO_SENDER_EMAIL','test@example.test'),patch.object(email_service.requests,'post',return_value=response) as post:
            self.assertEqual(email_service.enviar_reporte('actor@example.test','Reporte','contenido',[('reporte.pdf',b'PDF')],'stable'),'fake-id')
            self.assertEqual(post.call_args.kwargs['json']['attachment'][0]['content'],'UERG')
            post.side_effect=requests.Timeout()
            with self.assertRaises(email_service.ReportEmailError) as exc: email_service.enviar_reporte('a','s','c',[],'x')
            self.assertTrue(exc.exception.uncertain)
            post.side_effect=None;response.status_code=429
            with self.assertRaises(email_service.ReportEmailError) as exc: email_service.enviar_reporte('a','s','c',[],'x')
            self.assertTrue(exc.exception.retryable)
            self.assertFalse(exc.exception.uncertain)

    def test_existing_recovery_sender_remains_compatible(self):
        response = MagicMock(status_code=201)
        with patch.object(Config,'BREVO_API_KEY','fake'),patch.object(Config,'BREVO_SENDER_EMAIL','test@example.test'),patch.object(email_service.requests,'post',return_value=response):
            self.assertIsNone(email_service.enviar_codigo_recuperacion('test@example.test','123456'))


if __name__=='__main__':
    unittest.main()
